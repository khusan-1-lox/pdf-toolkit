import tempfile
import os
import shutil
import pytest
from docx import Document


def test_office_service_import():
    from app.services.office_convert import convert_office_to_pdf, convert_pdf_to_docx
    assert callable(convert_office_to_pdf)
    assert callable(convert_pdf_to_docx)


@pytest.fixture
def tmp_docx():
    d = Document()
    d.add_paragraph("Hello Word to PDF")
    p = tempfile.mktemp(suffix=".docx")
    d.save(p)
    yield p
    if os.path.exists(p):
        os.remove(p)


@pytest.fixture
def tmp_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello Excel to PDF"
    p = tempfile.mktemp(suffix=".xlsx")
    wb.save(p)
    yield p
    if os.path.exists(p):
        os.remove(p)


@pytest.fixture
def tmp_pptx():
    from pptx import Presentation

    prs = Presentation()
    slide_layout = prs.slide_layouts[0] if len(prs.slide_layouts) > 0 else prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    # try to set title if placeholder exists
    try:
        slide.shapes.title.text = "Hello PPT to PDF"
    except Exception:
        # fallback: add textbox
        txBox = slide.shapes.add_textbox(72, 72, 400, 100)
        txBox.text_frame.text = "Hello PPT to PDF"
    p = tempfile.mktemp(suffix=".pptx")
    prs.save(p)
    yield p
    if os.path.exists(p):
        os.remove(p)


def test_word_to_pdf_route(client, tmp_docx, monkeypatch):
    # Fallback if LibreOffice not installed: mock convert_office_to_pdf to create dummy PDF via PyMuPDF
    if shutil.which("soffice") is None:
        import fitz

        async def fake_convert(input_path: str, out_dir: str, timeout: int = 30) -> str:
            out_pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf")
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "dummy")
            doc.save(out_pdf)
            doc.close()
            return out_pdf

        monkeypatch.setattr("app.routers.word_to_pdf.convert_office_to_pdf", fake_convert)
        # also patch service directly in case of different import path
        monkeypatch.setattr("app.services.office_convert.convert_office_to_pdf", fake_convert)

    with open(tmp_docx, "rb") as f:
        resp = client.post(
            "/api/word-to-pdf",
            files={"file": ("test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    # verify response is a valid PDF (starts with %PDF)
    assert resp.content[:4] == b"%PDF"


def test_excel_to_pdf_route(client, tmp_xlsx, monkeypatch):
    if shutil.which("soffice") is None:
        import fitz

        async def fake_convert(input_path: str, out_dir: str, timeout: int = 30) -> str:
            out_pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf")
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "dummy")
            doc.save(out_pdf)
            doc.close()
            return out_pdf

        monkeypatch.setattr("app.routers.excel_to_pdf.convert_office_to_pdf", fake_convert)
        monkeypatch.setattr("app.services.office_convert.convert_office_to_pdf", fake_convert)

    with open(tmp_xlsx, "rb") as f:
        resp = client.post(
            "/api/excel-to-pdf",
            files={"file": ("test.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content[:4] == b"%PDF"


def test_ppt_to_pdf_route(client, tmp_pptx, monkeypatch):
    if shutil.which("soffice") is None:
        import fitz

        async def fake_convert(input_path: str, out_dir: str, timeout: int = 30) -> str:
            out_pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf")
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "dummy")
            doc.save(out_pdf)
            doc.close()
            return out_pdf

        monkeypatch.setattr("app.routers.ppt_to_pdf.convert_office_to_pdf", fake_convert)
        monkeypatch.setattr("app.services.office_convert.convert_office_to_pdf", fake_convert)

    with open(tmp_pptx, "rb") as f:
        resp = client.post(
            "/api/ppt-to-pdf",
            files={"file": ("test.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content[:4] == b"%PDF"


def test_pdf_to_word_route(client):
    import fitz
    from docx import Document

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello Word from PDF")
    pdf_bytes = doc.tobytes()
    doc.close()

    resp = client.post(
        "/api/pdf-to-word",
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    tmp = tempfile.mktemp(suffix=".docx")
    try:
        with open(tmp, "wb") as f:
            f.write(resp.content)
        d = Document(tmp)
        text = "\n".join([p.text for p in d.paragraphs])
        assert "Hello Word" in text or len(text.strip()) > 0, "Docx contains no data"
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def test_pdf_to_excel_route(client):
    import fitz
    import io
    from openpyxl import load_workbook

    # create tmp_pdf via fitz with text
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello Excel from PDF")
    page.insert_text((72, 90), "Second line")
    pdf_bytes = doc.tobytes()
    doc.close()

    resp = client.post(
        "/api/pdf-to-excel",
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # verify output can be opened via openpyxl and contains data
    tmp = tempfile.mktemp(suffix=".xlsx")
    try:
        with open(tmp, "wb") as f:
            f.write(resp.content)
        wb = load_workbook(tmp)
        ws = wb.active
        # check at least one cell has data containing expected text
        found = False
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and "Hello Excel" in str(cell):
                    found = True
                    break
            if found:
                break
        # fallback: any data at all
        if not found:
            has_any = any(cell is not None and str(cell).strip() != "" for row in ws.iter_rows(values_only=True) for cell in row if cell is not None)
            assert has_any, "Excel sheet contains no data"
        else:
            assert found
        wb.close()
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def test_pdf_to_ppt_route(client):
    import fitz
    from pptx import Presentation

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PPT from PDF")
    pdf_bytes = doc.tobytes()
    doc.close()

    resp = client.post(
        "/api/pdf-to-ppt",
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    tmp = tempfile.mktemp(suffix=".pptx")
    try:
        with open(tmp, "wb") as f:
            f.write(resp.content)
        prs = Presentation(tmp)
        # verify contains at least one slide with data (shapes)
        assert len(prs.slides) >= 1, "PPT should contain at least one slide"
        has_data = False
        for slide in prs.slides:
            if len(slide.shapes) > 0:
                has_data = True
                break
        assert has_data, "PPT slides contain no shapes"
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
